"""Safe text extraction for books, notes, transcripts, audio and video."""
from __future__ import annotations

import asyncio
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from app.config import get_settings

settings = get_settings()


@dataclass(slots=True)
class ExtractionResult:
    text: str
    method: str
    status: str = "ready"
    error: str | None = None


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".log"}
SUBTITLE_EXTENSIONS = {".srt", ".vtt"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"}


class _HTMLTextParser(HTMLParser):
    _SKIP = {"script", "style", "noscript", "svg"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag.lower() in self._SKIP:
            self._skip_depth += 1
        elif tag.lower() in {
            "p", "div", "br", "li", "h1", "h2", "h3", "h4", "section"
        }:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag.lower() in {
            "p", "div", "li", "h1", "h2", "h3", "h4", "section"
        }:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"[Page {index + 1}]\n{page_text}")
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        total_size = sum(info.file_size for info in archive.infolist())
        if total_size > settings.knowledge_max_extracted_bytes:
            raise ValueError("DOCX extracted content exceeds the configured safety limit.")
    from docx import Document

    document = Document(str(path))
    sections = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                sections.append(" | ".join(values))
    return "\n".join(sections)


def _extract_epub(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = [
            name for name in archive.namelist()
            if name.lower().endswith((".xhtml", ".html", ".htm"))
        ]
        total_size = sum(archive.getinfo(name).file_size for name in names)
        if total_size > settings.knowledge_max_extracted_bytes:
            raise ValueError("EPUB extracted content exceeds the configured safety limit.")
        sections: list[str] = []
        for name in sorted(names):
            parser = _HTMLTextParser()
            parser.feed(archive.read(name).decode("utf-8", errors="ignore"))
            value = _normalize("".join(parser.parts))
            if value:
                sections.append(value)
    return "\n\n".join(sections)


def _extract_subtitles(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"WEBVTT[^\n]*\n", "", raw, flags=re.IGNORECASE)
    raw = re.sub(r"^\d+\s*$", "", raw, flags=re.MULTILINE)
    raw = re.sub(
        r"\d{2}:\d{2}(?::\d{2})?[.,]\d{3}\s+-->.*$",
        "", raw, flags=re.MULTILINE,
    )
    raw = re.sub(r"<[^>]+>", " ", raw)
    return _normalize(raw)


def _run_ffmpeg(args: list[str]) -> subprocess.CompletedProcess[str] | None:
    if shutil.which("ffmpeg") is None:
        return None
    try:
        return subprocess.run(
            ["ffmpeg", *args],
            capture_output=True,
            text=True,
            timeout=settings.knowledge_media_command_timeout_seconds,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None


def _embedded_subtitles(path: Path) -> str:
    result = _run_ffmpeg([
        "-hide_banner", "-loglevel", "error", "-i", str(path),
        "-map", "0:s:0?", "-f", "webvtt", "-",
    ])
    if result is None or result.returncode != 0:
        return ""
    return _normalize(result.stdout)


def _extract_audio_for_transcription(path: Path) -> Path | None:
    temp_dir = Path(tempfile.mkdtemp(prefix="pamasmma-knowledge-audio-"))
    output = temp_dir / "audio.mp3"
    result = _run_ffmpeg([
        "-hide_banner", "-loglevel", "error", "-i", str(path),
        "-map", "0:a:0?", "-vn", "-ac", "1", "-ar", "16000",
        "-b:a", "32k", "-y", str(output),
    ])
    if result is None or result.returncode != 0 or not output.exists():
        shutil.rmtree(temp_dir, ignore_errors=True)
        return None
    return output


async def _transcribe(path: Path) -> str:
    if not settings.openai_api_key:
        return ""
    from openai import AsyncOpenAI

    client = AsyncOpenAI(api_key=settings.openai_api_key)
    try:
        with path.open("rb") as handle:
            response = await client.audio.transcriptions.create(
                model=settings.transcription_model,
                file=handle,
            )
        return _normalize(getattr(response, "text", "") or "")
    finally:
        await client.close()


async def extract_content(
    path: Path,
    filename: str,
    transcript_override: str | None = None,
) -> ExtractionResult:
    if transcript_override and transcript_override.strip():
        return ExtractionResult(_normalize(transcript_override), "user_supplied_transcript")

    extension = Path(filename).suffix.lower()
    try:
        if extension in TEXT_EXTENSIONS:
            text = path.read_text(encoding="utf-8", errors="ignore")
            return ExtractionResult(_normalize(text), "plain_text")
        if extension == ".pdf":
            return ExtractionResult(
                _normalize(await asyncio.to_thread(_extract_pdf, path)), "pypdf"
            )
        if extension == ".docx":
            return ExtractionResult(
                _normalize(await asyncio.to_thread(_extract_docx, path)), "python_docx"
            )
        if extension == ".epub":
            return ExtractionResult(
                _normalize(await asyncio.to_thread(_extract_epub, path)), "epub_html"
            )
        if extension in SUBTITLE_EXTENSIONS:
            return ExtractionResult(_extract_subtitles(path), "subtitle_file")

        if extension in VIDEO_EXTENSIONS:
            subtitles = await asyncio.to_thread(_embedded_subtitles, path)
            if subtitles:
                return ExtractionResult(subtitles, "embedded_subtitles")
            audio_path = await asyncio.to_thread(_extract_audio_for_transcription, path)
            if audio_path is None:
                return ExtractionResult(
                    "", "video_pending_transcription",
                    status="awaiting_transcription",
                    error="Video audio extraction is unavailable or no usable audio track exists.",
                )
            try:
                if audio_path.stat().st_size > settings.knowledge_transcription_max_bytes:
                    return ExtractionResult(
                        "", "video_transcription_limit",
                        status="awaiting_transcription",
                        error="Extracted audio exceeds the configured transcription size limit.",
                    )
                transcript = await _transcribe(audio_path)
            finally:
                shutil.rmtree(audio_path.parent, ignore_errors=True)
            if transcript:
                return ExtractionResult(transcript, "video_ffmpeg_transcription")
            return ExtractionResult(
                "", "video_pending_transcription",
                status="awaiting_transcription",
                error="No transcription provider is configured. Add an OpenAI API key or upload a transcript.",
            )

        if extension in AUDIO_EXTENSIONS:
            if path.stat().st_size > settings.knowledge_transcription_max_bytes:
                return ExtractionResult(
                    "", "audio_transcription_limit",
                    status="awaiting_transcription",
                    error="Audio exceeds the configured transcription size limit.",
                )
            transcript = await _transcribe(path)
            if transcript:
                return ExtractionResult(transcript, "audio_transcription")
            return ExtractionResult(
                "", "audio_pending_transcription",
                status="awaiting_transcription",
                error="No transcription provider is configured. Add an OpenAI API key.",
            )

        return ExtractionResult(
            "", "unsupported", status="failed",
            error=f"Unsupported file type: {extension or 'unknown'}",
        )
    except Exception as exc:
        return ExtractionResult("", "failed", status="failed", error=str(exc)[:500])
